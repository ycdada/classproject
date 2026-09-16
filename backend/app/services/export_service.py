"""
Export exams to Word (.docx) and PDF formats.
Supports with/without answer variants.
"""
import io
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from ..models.exam import Exam

TYPE_NAMES = {
    "choice": "选择题",
    "fill": "填空题",
    "tf": "判断题",
    "short_answer": "简答题",
    "code": "编程题",
}


class ExportService:
    def export_word(self, exam: Exam, with_answer: bool = False) -> bytes:
        """Export exam to Word (.docx) format."""
        doc = Document()

        # Title
        title = doc.add_heading(exam.title, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Info line
        info = doc.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        info.add_run(f"总分: {exam.total_score}分 | 考试时间: {exam.duration}分钟").font.size = Pt(10)

        doc.add_paragraph()  # spacer

        # Group questions by type
        grouped = {}
        for eq in sorted(exam.questions, key=lambda x: x.sort_order):
            q = eq.question
            qtype = q.type if hasattr(q, 'type') else "choice"
            if qtype not in grouped:
                grouped[qtype] = []
            grouped[qtype].append(eq)

        # Question numbering
        q_num = 1
        for qtype in ["choice", "fill", "tf", "short_answer", "code"]:
            if qtype not in grouped:
                continue
            type_name = TYPE_NAMES.get(qtype, qtype)
            doc.add_heading(f"{'一' if qtype == 'choice' else '二' if qtype == 'fill' else '三' if qtype == 'tf' else '四' if qtype == 'short_answer' else '五'}、{type_name}", level=2)

            for eq in grouped[qtype]:
                q = eq.question
                # Question content
                p = doc.add_paragraph()
                p.add_run(f"{q_num}. ").bold = True
                p.add_run(f"({eq.score}分) ").font.size = Pt(9)
                content = q.content if hasattr(q, 'content') else str(q)
                p.add_run(content)

                # Options for choice questions
                if qtype == "choice":
                    options = q.options if hasattr(q, 'options') else None
                    if options:
                        for opt in options:
                            doc.add_paragraph(f"    {opt}")

                # Answer and explanation (only if with_answer)
                if with_answer:
                    answer_p = doc.add_paragraph()
                    answer = q.answer if hasattr(q, 'answer') else ""
                    answer_p.add_run(f"答案: {answer}").font.color.rgb = RGBColor(0, 128, 0)

                    explanation = q.explanation if hasattr(q, 'explanation') and q.explanation else ""
                    if explanation:
                        exp_p = doc.add_paragraph()
                        exp_p.add_run(f"解析: {explanation}").font.size = Pt(9)

                doc.add_paragraph()  # spacer
                q_num += 1

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue()

    def export_pdf(self, exam: Exam, with_answer: bool = False) -> bytes:
        """Export exam to PDF format."""
        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=A4,
            leftMargin=2*cm, rightMargin=2*cm,
            topMargin=2*cm, bottomMargin=2*cm,
        )

        styles = getSampleStyleSheet()
        story = []

        # Title
        title_style = ParagraphStyle(
            'CustomTitle', parent=styles['Title'],
            fontSize=18, spaceAfter=12, alignment=1,
            fontName='Helvetica-Bold',
        )
        story.append(Paragraph(exam.title, title_style))

        # Info
        info_style = ParagraphStyle(
            'Info', parent=styles['Normal'],
            fontSize=10, alignment=1, spaceAfter=20,
        )
        story.append(Paragraph(
            f"总分: {exam.total_score}分 | 考试时间: {exam.duration}分钟",
            info_style,
        ))

        # Group questions by type
        grouped = {}
        for eq in sorted(exam.questions, key=lambda x: x.sort_order):
            q = eq.question
            qtype = q.type if hasattr(q, 'type') else "choice"
            if qtype not in grouped:
                grouped[qtype] = []
            grouped[qtype].append(eq)

        q_num = 1
        for qtype in ["choice", "fill", "tf", "short_answer", "code"]:
            if qtype not in grouped:
                continue

            type_name = TYPE_NAMES.get(qtype, qtype)
            heading_style = ParagraphStyle(
                'Heading', parent=styles['Heading2'],
                fontSize=14, spaceAfter=8, spaceBefore=12,
            )
            story.append(Paragraph(type_name, heading_style))

            for eq in grouped[qtype]:
                q = eq.question
                content = q.content if hasattr(q, 'content') else str(q)

                q_style = ParagraphStyle(
                    'Question', parent=styles['Normal'],
                    fontSize=11, spaceAfter=6, leading=16,
                )
                story.append(Paragraph(
                    f"<b>{q_num}.</b> ({eq.score}分) {content}",
                    q_style,
                ))

                # Options for choice
                if qtype == "choice":
                    options = q.options if hasattr(q, 'options') else None
                    if options:
                        for opt in options:
                            story.append(Paragraph(
                                f"&nbsp;&nbsp;&nbsp;&nbsp;{opt}",
                                styles['Normal'],
                            ))

                if with_answer:
                    answer = q.answer if hasattr(q, 'answer') else ""
                    ans_style = ParagraphStyle(
                        'Answer', parent=styles['Normal'],
                        fontSize=10, textColor=colors.green,
                    )
                    story.append(Paragraph(f"<b>答案:</b> {answer}", ans_style))

                    explanation = q.explanation if hasattr(q, 'explanation') and q.explanation else ""
                    if explanation:
                        story.append(Paragraph(f"<i>解析: {explanation}</i>", styles['Normal']))

                story.append(Spacer(1, 6))
                q_num += 1

        doc.build(story)
        buf.seek(0)
        return buf.getvalue()
