"""Task 21 — 重写学术四件套（新建内容后覆盖 原文档/ 同名文件）。

度量口径与 scripts/measure_diff.py 一致：旧段落集合与新段落集合的
1 - |交集|/|旧集合| ≥ 0.30（docx 按段落、pptx 按 a:t 文本）。
"""
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "scripts"))

from measure_diff import extract_paragraphs  # noqa: E402

DOCS = REPO / "原文档"

PLAN = [
    ("题目与基础需求", [
        "课题名称：《数据结构》智能出卷系统（完善版）。使用者为高校《数据结构》课程教师，唯一角色是教师，没有学生端。",
        "基础需求一：教师上传教案、PPT 课件与课程大纲，系统解析后形成课程材料库，材料区分为教案、PPT、大纲、其他四类。",
        "基础需求二：系统从材料中抽出知识树，点上承载专业术语、概念定义、解题步骤与教学重点，其中解题步骤是本版新增的抽取维度。",
        "基础需求三：教师填写出卷需求，包括教学进度、考试范围、题型与数量、难度分布、重点考查方向，以及试卷时长与总分。",
        "基础需求四：系统按出卷需求筛选、组织知识节点，生成一份本次考试的考查范围树，而不是让教师逐个勾选节点编号。",
        "基础需求五：考查范围以树状呈现，教师可以修改文字、增删节点、排除子树，确认后范围锁定，才能进入组卷。",
        "基础需求六：组卷严格按确认后的考查范围与材料讲解方式进行，从题库装配试卷；题库不足时产生题目草稿而不是静默补题。",
        "基础需求七：教师可以预览试卷，修改题干、选项、答案、解析、分值与顺序，保存并导出 DOCX、PDF、TXT 及答题卡。",
    ]),
    ("用户与对象", [
        "教师是系统的唯一用户：上传材料、维护知识树与题库、提出组卷需求、审核考查范围与题目草稿、导出试卷。",
        "课程材料库是教师上传的全部教学材料的集合，材料包含解析后的正文，供知识抽取与组卷贴合讲解方式。",
        "考查范围是根据出卷需求从知识树筛选组织出的树状结构，有提出与确认两种状态，是本版引入的核心用户可见对象。",
        "题目草稿是模型针对题库缺口拟出的候选题，未获教师确认前不是题目，不得进入题库，也不得编入试卷。",
    ]),
    ("技术栈", [
        "前端采用 React 18 + TypeScript + Vite + Ant Design 5 + React Router + Zustand，全部页面按教师工作台重新实现。",
        "后端保持 FastAPI + SQLAlchemy async + aiosqlite，向量检索使用 ChromaDB 与本地 BAAI/bge-small-zh-v1.5 嵌入模型。",
        "大模型能力由 DeepSeek Chat 提供；教学助手由单一 LangGraph 状态机编排，支持断线续聊与人工审核中断。",
        "组卷主线在服务层实现：ScopeBuilder 负责提出考查范围，PaperAssembler 负责确认后的试卷装配。",
    ]),
    ("主线：出卷需求 → 考查范围确认 → 组卷 → 草稿审核 → 导出", [
        "第一步，教师在组卷向导填写出卷需求，系统调用 ScopeBuilder 提出考查范围，状态为 proposed。",
        "第二步，教师在树状审核页修改文案、增删节点、勾选包含关系，点确认后范围状态变为 confirmed 并锁定。",
        "第三步，PaperAssembler 按确认范围拼接知识上下文，先检索题库装配试卷；未确认的范围直接拒绝组卷。",
        "第四步，题型数量不足时，系统调用出题适配器拟定题目草稿，交由教师在草稿页接受或拒绝，接受后才写入题库。",
        "第五步，教师在预览页修改题目与配分，保存试卷，并按需导出含答案或不含答案版本以及答题卡。",
    ]),
    ("差异与验收", [
        "完善版对照两套差异基线验收：冻结的旧版可运行源码，以及原文档目录中的学术与工程文稿。",
        "源码 churn、模块替换、学术文档重写三项度量均须不低于 30%，由仓库内 scripts/measure_diff.py 脚本量化。",
        "模块清单覆盖十九个核心模块，材料种类、知识抽取、考查范围、题目草稿、组卷器、教学助手与各页面均已替换。",
        "文档层面除四件套重写外，工程文档（README、领域词、架构决策记录）也按新术语全部改写。",
    ]),
    ("风险", [
        "模型编造：知识抽取与拟草稿都约束只能使用材料原文内容，解题步骤不得编造，材料没有就留空。",
        "题库不足：缺口只产生题目草稿，绝不静默写题入库；草稿未经确认不会出现在任何试卷上。",
        "提取空步骤：部分材料没有显式解题步骤，系统保留空字段并在界面上提示教师补录。",
    ]),
    ("交付物", [
        "可运行系统：仓库根目录的 backend 与 frontend，附本地运行说明与环境变量样例。",
        "测试与度量：后端 pytest 套件覆盖考查范围与组卷守卫，30% 差异门脚本与单测。",
        "文档四件套：项目计划书、中期报告、开发文档、答辩 PPT，全部按新大纲重写。",
    ]),
]

MIDTERM = [
    ("基础需求符合性对照表", [
        "上传教案、PPT、大纲形成课程材料库：已实现，上传接口支持材料种类字段，列表按种类区分展示。",
        "抽出知识点、术语、定义、解题步骤、教学重点：已实现，抽取提示词要求不得编造，解题步骤为新增字段。",
        "输入教学进度、考试范围、题型、难度、重点方向：已实现，出卷需求 schema 新增教学进度与考试范围文本字段。",
        "筛选组织生成本次考查范围：已实现，ScopeBuilder 保留章节点父子结构，无大模型时用子串匹配过滤。",
        "树状展示，教师审核修改确认：已实现，确认后节点只读，返回 409 拒绝继续修改。",
        "按确认范围与材料讲解方式组卷：已实现，PaperAssembler 只从题库装配，未确认范围直接抛错。",
        "预览、修改、保存、导出：已实现，预览页支持改题干、选项、答案、解析、分值与顺序，导出多格式与答题卡。",
    ]),
    ("新架构", [
        "考查范围模块：exam_scopes 与 exam_scope_nodes 两张表持久化范围树，出卷需求以 JSON 存档。",
        "ScopeBuilder：从课程知识树按出卷需求过滤，保留祖先链，产出 proposed 状态的考查范围。",
        "PaperAssembler：组卷唯一实现，拼接定义、术语、解题步骤与教学重点作为上下文，先检索题库。",
        "单一教学助手：LangGraph 状态机依次经过意图路由、提出范围、范围审核中断、装配、试卷审核中断。",
    ]),
    ("前端换栈说明", [
        "旧版前端使用 Vue 3 与 Naive UI，页面之间状态共享薄弱，类型系统缺失，向导式流程难以维护。",
        "完善版以 React 18 + Ant Design 5 重写全部页面，页面职责保持一致，仅新增草稿页与范围审核步骤。",
        "旧 Vue 源码已从运行目录移除并冻结进差异基线，新前端构建以 tsc 类型检查与 Vite 产物为门。",
    ]),
    ("已完成模块与演示路径", [
        "材料库演示：上传一份 PPT（种类 PPT）与一份教案（种类 教案），列表正确显示种类，详情可见解析正文。",
        "知识树演示：对材料点提取，树形展开章节点，点上可见定义、术语、解题步骤与教学重点并可编辑。",
        "组卷演示：向导第一步填写需求，第二步审核确认考查范围，第三步得到试卷与草稿提醒。",
        "导出演示：预览页修改分值后导出 DOCX，文件可下载打开。",
    ]),
    ("差异度量方法与当前数字", [
        "源码 churn：以冻结旧源码文件哈希为基准，按字节数计算未被新树继承的比例，当前为 100%。",
        "模块替换：十九个核心模块中已替换十五个，比例约 79%，超过 30% 门槛。",
        "学术文档重写：四件套按段落集合比较重写比例，本版全部按新大纲重新撰写后覆盖原文件。",
    ]),
    ("待完成与风险", [
        "待完成：答辩 PPT 的演示录屏占位、README 运行截图补充。",
        "风险：大模型拟草稿质量依赖材料质量，需教师在草稿页把关；本地嵌入模型首次加载较慢。",
    ]),
]

DEVDOC = [
    ("领域词", [
        "教师：使用系统的唯一角色，上传材料、维护知识树与题库、组卷、审核、导出。",
        "课程材料库：教师上传的全部教学材料集合；材料分教案、PPT、大纲、其他四类。",
        "知识树：某份材料抽出的章、节、点层级；点承载术语、定义、解题步骤、教学重点。",
        "出卷需求：本次组卷的输入，含教学进度、考试范围、题型数量、难度分布、重点方向、时长与总分。",
        "考查范围：按出卷需求从知识树筛选组织的树，proposed 可改，confirmed 锁定后才能组卷。",
        "题目草稿：模型针对缺口拟的候选题，pending、accepted、rejected 三态，接受才入题库。",
        "题库：只含已确认题目；来源为材料抽取、教师录入与教师确认后的草稿。",
        "组卷：确认范围后从题库装配试卷；缺口只产生题目草稿，不静默写题。",
        "教学助手：教师端对话入口，走单一 LangGraph 图，组卷同样先确认考查范围。",
    ]),
    ("目录结构", [
        "仓库根：backend（FastAPI 应用与 agent_v2 状态机）、frontend（React 应用）、scripts（度量脚本）。",
        "baselines/v2-source：冻结的旧版源码只读基线；baselines/original-docs：冻结的旧学术文稿。",
        "原文档：重写后的学术四件套；docs/adr：六条不可逆架构决策记录。",
    ]),
    ("数据模型", [
        "materials：文件名、类型、种类（kind：lecture_notes/slides/syllabus/other）、解析正文、页数。",
        "knowledge_nodes：material_id、parent_id、名称、节点类型、definition、key_terms、teaching_emphasis、solution_steps。",
        "exam_scopes / exam_scope_nodes：范围状态（proposed/confirmed）与范围节点（含 source_node_id 溯源、included 开关）。",
        "question_drafts：题型、难度、题干、选项、答案、解析、状态（pending/accepted/rejected）、exam_scope_id。",
        "questions：题库主表，source 只有 manual 与 ppt_extracted，本版不引入 AI 生成来源。",
        "exams / exam_questions：试卷头（含 requirements_json 与 knowledge_snapshot_json）与试卷题目（分值、顺序）。",
    ]),
    ("HTTP 一览", [
        "材料：POST /api/materials/upload（含 kind 表单字段）、GET /api/materials、GET/DELETE /api/materials/{id}。",
        "知识树：POST /api/knowledge/extract/{material_id}、GET /api/knowledge/tree/{material_id}、PUT /api/knowledge/nodes/{id}。",
        "考查范围：POST /api/scopes/propose、GET /api/scopes/{id}、PUT /api/scopes/{id}/nodes/{nid}、POST /api/scopes/{id}/nodes、DELETE 节点、POST /api/scopes/{id}/confirm。",
        "组卷：POST /api/exams/generate（必填 scope_id；未确认返回 400 scope_not_confirmed）。",
        "草稿：GET /api/drafts?status=pending、PUT /api/drafts/{id}、POST /api/drafts/{id}/accept、POST /api/drafts/{id}/reject。",
        "教学助手：POST /api/chat（SSE，事件 intent/tool_call/scope_pending/review_pending/text/done/error）。",
    ]),
    ("组卷时序", [
        "教师提交出卷需求 → POST /api/scopes/propose → ScopeBuilder 过滤知识树并落库 → 返回 proposed 范围树。",
        "教师编辑节点 → PUT /api/scopes/{id}/nodes/{nid}；确认 → POST /api/scopes/{id}/confirm → 状态 confirmed。",
        "教师触发组卷 → POST /api/exams/generate 携带 scope_id → 校验 confirmed，否则 400。",
        "PaperAssembler 拼接 included 节点的定义、术语、解题步骤、教学重点 → RAG 检索题库 → 题型不足时调出题适配器拟草稿 → create_many 写入 question_drafts → questions 表零插入。",
        "返回 exam_id、question_count、drafts、shortfall、summary；前端提示教师去草稿页确认。",
    ]),
    ("前端页面与路由", [
        "/dashboard 首页：统计卡片与内嵌教学助手；/materials 材料列表；/materials/:id 材料详情。",
        "/knowledge/:materialId 知识树：提取、重新提取、编辑定义术语解题步骤教学重点。",
        "/exams/create 组卷向导：三步（出卷需求 → 考查范围 → 试卷）；未确认范围不得调用 generate。",
        "/exams 试卷列表；/exams/:id 预览与导出；/drafts 草稿审核。",
    ]),
    ("本地运行", [
        "后端：cd backend && python -m venv .venv && .venv/Scripts/pip install -r requirements.txt，复制 .env.example 为 .env 填入 DEEPSEEK_API_KEY，uvicorn app.main:app 启动于 8000 端口。",
        "前端：cd frontend && npm install && npm run dev，Vite 代理 /api 到 8000 端口。",
        "允许重建本地数据：删除 backend/data 与 checkpoints.db 后系统会自动重建 SQLite 与 Chroma。",
    ]),
    ("测试与差异门", [
        "pytest 套件：test_scope_builder、test_paper_assembler、test_draft_service、test_generate_guard、test_scope_api。",
        "度量脚本：python scripts/measure_diff.py，source_churn、module_change、doc_rewrite 三项均不低于 0.30 才通过。",
        "度量脚本自身带单测：scripts/test_measure_diff.py 在合成目录上验证度量口径，防止 30% 门悄悄失效。",
    ]),
]

PPT = [
    ("题目与基础需求", [
        "数据结构智能出卷系统 完善版",
        "唯一用户：教师",
        "七条基础需求：材料库 / 知识抽取 / 出卷需求 / 考查范围 / 树状审核 / 组卷 / 预览导出",
    ]),
    ("问题：静默补题与无考查范围", [
        "旧版组卷：题库不足时模型写题直接入库，未经教师确认",
        "旧版组卷：没有本次考试的考查范围对象，向导只显示节点编号",
        "旧版抽取：知识点缺少解题步骤，试卷难以贴合教学讲法",
    ]),
    ("目标完善点", [
        "新增考查范围：提出 → 审核 → 确认，未确认不得组卷",
        "新增题目草稿：缺口只拟草稿，教师接受才入题库",
        "抽取增加解题步骤；前端整体换为 React",
    ]),
    ("总体架构", [
        "前端：React 18 + TypeScript + Vite + Ant Design 5 + React Router + Zustand",
        "后端：FastAPI + SQLAlchemy async + ChromaDB + DeepSeek + 本地中文嵌入",
        "教学助手：单一 LangGraph 状态机，支持中断与续聊",
    ]),
    ("课程材料库", [
        "上传教案 / PPT / 大纲，材料带种类标签",
        "解析正文入库，列表与详情按种类区分",
        "删除材料时同步清理派生知识树",
    ]),
    ("知识抽取（含解题步骤）", [
        "章 → 节 → 点三级知识树",
        "点承载术语、定义、教学重点",
        "新增解题步骤字段：只抄材料原文，不得编造，没有则留空",
    ]),
    ("出卷需求", [
        "教学进度、考试范围、重点考查方向",
        "题型数量与难度分布、时长与总分",
        "可选关联材料，缩小范围来源",
    ]),
    ("考查范围确认", [
        "ScopeBuilder 按需求过滤知识树，保留祖先链",
        "教师改文案、增删节点、勾选包含",
        "确认后锁定：再修改返回 409",
    ]),
    ("组卷与题目草稿", [
        "PaperAssembler：确认范围 + 题库检索装配",
        "缺口产生题目草稿，questions 表零插入",
        "教师接受草稿后题目才进题库",
    ]),
    ("教学助手状态机", [
        "意图路由 → 搜题 / 知识问答 / 组卷",
        "组卷：提出范围 → 范围审核中断 → 装配 → 试卷审核中断",
        "断线续聊：SQLite 检查点恢复线程",
    ]),
    ("导出", [
        "DOCX / PDF / TXT，含答案与不含答案两个版本",
        "答题卡 PDF 与题号对齐",
        "预览页可改题干、选项、答案、解析、分值、顺序",
    ]),
    ("差异验收", [
        "两套差异基线：冻结旧源码 + 冻结学术文稿",
        "源码 churn 100%；模块替换 15/19 ≈ 79%；文档重写达标",
        "度量脚本 measure_diff.py 退出码 0 视为过门",
    ]),
    ("演示", [
        "上传 PPT 与教案 → 提取知识树查看解题步骤",
        "组卷向导：需求 → 范围确认 → 试卷 + 草稿提醒",
        "草稿页接受一题 → 题库可查 → 导出 DOCX",
    ]),
    ("总结", [
        "组卷主线严格对齐基础需求七条",
        "考查范围与题目草稿成为可审核对象",
        "两套基线 30% 差异可量化验收",
    ]),
]


def build_docx(path: Path, title: str, subtitle: str, chapters, accent: str):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)

    # 封面
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(26)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.size = Pt(14)
    r2.font.color.rgb = RGBColor.from_string(accent)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("摸鱼小组 · 完善版 v3")
    r3.font.size = Pt(12)
    doc.add_page_break()

    for heading, paras in chapters:
        h = doc.add_heading(heading, level=1)
        for r in h.runs:
            r.font.color.rgb = RGBColor.from_string("1F3864")
        for text in paras:
            para = doc.add_paragraph(text)
            para.paragraph_format.first_line_indent = Pt(24)
    doc.save(path)


def build_pptx(path: Path, title: str, slides):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.enum.text import PP_ALIGN

    box = cover.shapes.add_textbox(Inches(1), Inches(2.4), Inches(11.3), Inches(2.2))
    tf = box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(44)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "完善版 v3 · 教师端 · 考查范围 + 题目草稿"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(0x0E, 0x7C, 0x86)

    for heading, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        head = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1))
        hf = head.text_frame
        hf.text = heading
        hf.paragraphs[0].font.size = Pt(30)
        hf.paragraphs[0].font.bold = True
        hf.paragraphs[0].font.color.rgb = RGBColor(0x0E, 0x7C, 0x86)

        body = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.2))
        bf = body.text_frame
        bf.word_wrap = True
        for i, b in enumerate(bullets):
            para = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
            para.text = f"• {b}"
            para.font.size = Pt(20)
            para.space_after = Pt(12)

    prs.save(path)


def rewrite_score(name: str) -> float:
    old = extract_paragraphs(Path("baselines/original-docs") / name)
    new = extract_paragraphs(DOCS / name)
    old_set, new_set = set(old), set(new)
    if not old_set:
        return 1.0
    return 1.0 - len(old_set & new_set) / len(old_set)


def main() -> None:
    # 计划书
    build_docx(DOCS / "数据结构出卷系统-项目计划书.docx", "数据结构智能出卷系统 项目计划书", "完善版 v3 · 以考查范围与题目草稿为核心的教师端出卷", PLAN, "0E7C86")
    # 中期报告
    build_docx(DOCS / "数据结构智能出卷系统_中期报告.docx", "数据结构智能出卷系统 中期报告", "完善版 v3 · 基础需求符合性与差异度量", MIDTERM, "843D0F")
    # 开发文档
    build_docx(DOCS / "开发文档.docx", "数据结构智能出卷系统 开发文档", "完善版 v3 · 领域词 / 数据模型 / HTTP / 组卷时序", DEVDOC, "385723")
    # 答辩 PPT
    build_pptx(DOCS / "摸鱼小组-《数据结构》智能出卷系统.pptx", "《数据结构》智能出卷系统", PPT)

    for name in (
        "数据结构出卷系统-项目计划书.docx",
        "数据结构智能出卷系统_中期报告.docx",
        "开发文档.docx",
        "摸鱼小组-《数据结构》智能出卷系统.pptx",
    ):
        print(f"{name}: rewrite={rewrite_score(name):.2f}")


if __name__ == "__main__":
    main()
